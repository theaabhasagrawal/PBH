"""
PB Health Hospital - real sales-call analysis exported to Excel.

    python export_call_analysis.py                       # last 21 days up to today, live portal API
    python export_call_analysis.py --days 7
    python export_call_analysis.py --from 2026-09-25 --to 2026-10-16
    python export_call_analysis.py --max-calls 300       # quick run
    python export_call_analysis.py --cached              # rebuild from transcripts already downloaded (no API)
    python export_call_analysis.py --from-db             # rebuild from data/pb_health.db (demo rows excluded)

Pipeline: search (keyword clusters + category 10693) -> full transcripts -> harmonize
(AGENT / CUSTOMER lines, Hinglish normalized into pitch tags) -> per-call facts -> workbook.

Outputs
    output/PB_Health_Call_Analysis.xlsx         4 tabs: Executive Summary, Customer Doubts,
                                                Pitch Elements, Call-by-Call Master Data
    output/PB_Health_Harmonized_Dialogue.csv    every analysed line, verbatim + harmonized
    data/transcripts/<call_id>.json             transcript cache (re-runs skip the API)

The portal (datascienceapi.policybazaar.com) only resolves on the PB office network / VPN.
"""
from __future__ import annotations

import argparse
import csv
import json
import logging
import math
import re
import sqlite3
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config
from call_facts import DOUBTS, ELEMENTS, GENERIC_GURGAON, HUBS, CallFacts, analyze_call, hospital_window

log = logging.getLogger("pb.export")

SEARCH_TERMS: dict[str, list[str]] = {
    "brand": ["pb health", "pb hospital", "policybazaar hospital", "pb health hospital"],
    "hinglish": ["khud ka hospital", "apna hospital", "humara hospital"],
    "hubs": ["iffco chowk", "sector 62", "golf course", "faridabad", "sector 2", "noida sector 50", "sector 50", "rohini"],
}
CATEGORY_LABEL = f"category {config.PB_HEALTH_CATEGORY_ID}"
CACHE_DIR = config.BASE_DIR / "data" / "transcripts"
OUT_XLSX = config.OUTPUT_DIR / "PB_Health_Call_Analysis.xlsx"
OUT_CSV = config.OUTPUT_DIR / "PB_Health_Harmonized_Dialogue.csv"


@dataclass
class Record:
    call_id: str
    meta: dict
    english: list[dict]
    hindi: list[dict]


@dataclass
class RunInfo:
    start: date
    end: date
    source: str                         # live API | transcript cache | SQLite DB
    scanned: int = 0                    # unique calls returned by search / loaded
    failed: int = 0                     # transcript fetch errors
    empty: int = 0                      # transcript fetched but no usable lines
    qualified_ncr: int | None = None    # portal count of all qualified NCR health sales calls
    qualified_days: int = 0
    notes: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------- #
# Ingestion
# --------------------------------------------------------------------------- #
def _meta(hit) -> dict:
    return {
        "call_id": hit.call_id, "call_date": hit.call_date, "agent_id": hit.agent_id, "agent_name": hit.agent_name,
        "city": hit.city, "state": hit.state, "talktime": hit.talktime, "lead_id": hit.lead_id,
        "matched_terms": sorted(hit.matched_terms),
    }


def collect_calls(client, start: date, end: date, include_category: bool = True) -> dict:
    """Run every search phrase (and the category-only query) over start..end. Returns {call_id: SearchHit}."""
    from api_client import SearchHit, iter_windows

    groups: list[tuple[list[str], str, list[int] | None]] = [
        ([t], t, None) for terms in SEARCH_TERMS.values() for t in terms]   # None -> api_client default (cat 10693)
    if include_category:
        groups.append(([], CATEGORY_LABEL, [config.PB_HEALTH_CATEGORY_ID]))
    hits: dict = {}
    for w_from, w_to in iter_windows(start, end):
        for terms, label, cats in groups:
            rows, _ = client.search_window(terms, w_from, w_to, categories=cats)
            for row in rows:
                h = SearchHit.from_row(row, term=label)
                if h is None:
                    continue
                if h.call_id in hits:
                    hits[h.call_id].matched_terms |= h.matched_terms
                else:
                    hits[h.call_id] = h
        log.info("Searched %s -> %s: %d unique calls so far", w_from.date(), w_to.date(), len(hits))
    return hits


def _cache_path(call_id: str) -> Path:
    return CACHE_DIR / f"{re.sub(r'[^A-Za-z0-9_.-]', '_', str(call_id))}.json"


def _load_cache(call_id: str) -> dict | None:
    p = _cache_path(call_id)
    if not p.exists():
        return None
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except ValueError:
        return None


def _save_cache(rec: Record) -> None:
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    _cache_path(rec.call_id).write_text(json.dumps(
        {"meta": rec.meta, "english": rec.english, "hindi": rec.hindi,
         "fetched_at": datetime.now().isoformat(timespec="seconds")}, ensure_ascii=False), encoding="utf-8")


def fetch_transcripts(client, hits: dict, workers: int) -> tuple[list[Record], int]:
    records: list[Record] = []
    todo = []
    for h in hits.values():
        cached = _load_cache(h.call_id)
        if cached and (cached.get("english") or cached.get("hindi")):
            rec = Record(h.call_id, _meta(h), cached.get("english") or [], cached.get("hindi") or [])
            _save_cache(rec)   # refresh metadata (matched terms) without re-fetching
            records.append(rec)
        else:
            todo.append(h)
    log.info("%d transcripts cached, %d to fetch", len(records), len(todo))
    failed = 0
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futures = {pool.submit(client.get_transcript, h.call_id): h for h in todo}
        for i, fut in enumerate(as_completed(futures), 1):
            h = futures[fut]
            try:
                t = fut.result()
            except Exception as exc:  # noqa: BLE001 - one bad call must not stop the run
                failed += 1
                log.warning("Transcript %s failed: %s", h.call_id, exc)
                continue
            rec = Record(h.call_id, _meta(h), t.get("english") or [], t.get("hindi") or [])
            _save_cache(rec)
            records.append(rec)
            if i % 50 == 0:
                log.info("Fetched %d/%d transcripts", i, len(todo))
    return records, failed


def load_from_live(start: date, end: date, max_calls: int | None, include_category: bool,
                   denominator: bool, workers: int) -> tuple[list[Record], RunInfo]:
    from api_client import PortalClient

    run = RunInfo(start, end, "Policybazaar DS portal - live API")
    client = PortalClient()
    client.login()
    hits = collect_calls(client, start, end, include_category)
    run.scanned = len(hits)
    if max_calls and len(hits) > max_calls:
        keep = sorted(hits, key=lambda c: hits[c].call_date or "", reverse=True)[:max_calls]
        hits = {c: hits[c] for c in keep}
        run.notes.append(f"--max-calls {max_calls}: analysed the {max_calls} most recent of {run.scanned} calls found.")
    records, run.failed = fetch_transcripts(client, hits, workers)
    if denominator:
        total, days = 0, 0
        d = start
        while d <= end:
            n, _ = client.count_qualified_ncr_calls(d)
            if n is not None:
                total, days = total + n, days + 1
            d += timedelta(days=1)
        n_days = (end - start).days + 1
        if days:
            run.qualified_ncr, run.qualified_days = total, days
            if days < n_days:
                run.notes.append(f"Qualified NCR call count available for {days} of {n_days} days only.")
    return records, run


def _in_range(call_date: str | None, start: date, end: date) -> bool:
    if not call_date:
        return True
    try:
        d = date.fromisoformat(str(call_date)[:10])
    except ValueError:
        return True
    return start <= d <= end


def load_from_cache(start: date, end: date) -> tuple[list[Record], RunInfo]:
    run = RunInfo(start, end, f"transcript cache ({CACHE_DIR})")
    records = []
    for p in sorted(CACHE_DIR.glob("*.json")):
        try:
            d = json.loads(p.read_text(encoding="utf-8"))
        except ValueError:
            continue
        meta = d.get("meta") or {}
        if _in_range(meta.get("call_date"), start, end):
            records.append(Record(str(meta.get("call_id") or p.stem), meta, d.get("english") or [], d.get("hindi") or []))
    run.scanned = len(records)
    return records, run


def load_from_db(db_path: Path, start: date, end: date) -> tuple[list[Record], RunInfo]:
    run = RunInfo(start, end, f"SQLite DB ({db_path}), demo rows excluded")
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute(
            """SELECT c.*, t.english AS t_en, t.hindi AS t_hi FROM calls c JOIN transcripts t USING(call_id)
               WHERE COALESCE(c.source, '') != 'demo' AND c.call_date BETWEEN ? AND ?""",
            (start.isoformat(), end.isoformat())).fetchall()
    finally:
        conn.close()
    records = []
    for r in rows:
        meta = {k: r[k] for k in ("call_id", "agent_id", "agent_name", "city", "state", "talktime", "lead_id")}
        meta["call_date"] = r["call_ts"] or r["call_date"]
        meta["matched_terms"] = json.loads(r["matched_terms"] or "[]")
        records.append(Record(str(r["call_id"]), meta, json.loads(r["t_en"] or "[]"), json.loads(r["t_hi"] or "[]")))
    run.scanned = len(records)
    return records, run


# --------------------------------------------------------------------------- #
# Tables
# --------------------------------------------------------------------------- #
def pct(n: int, d: int | None) -> float | None:
    return n / d if d else None


def _counter_scripts() -> dict[str, str]:
    from deck_generator import DECK

    idx = {"trust": 0, "location": 1, "cashless": 2, "scope": 3, "discount": 4, "skeptic": 5}
    out = {}
    for key, i in idx.items():
        _title, _cust, say, tip = DECK["objections"][i]
        out[key] = re.sub(r"<[^>]+>", "", f"Say: {say}\nTip: {tip}")
    return out


def _date_str(meta: dict) -> str:
    d = str(meta.get("call_date") or "")
    return d[:16].replace("T", " ") if len(d) >= 16 else d


def _minutes(secs: object) -> float | None:
    try:
        return round(float(secs) / 60, 1)
    except (TypeError, ValueError):
        return None


def _yn(c: CallFacts, key: str) -> str:
    return "N/A" if not c.facts.pitched else ("Yes" if c.facts.elements.get(key) else "No")


def _multi_value(c: CallFacts) -> str:
    if not c.facts.pitched:
        return "N/A"
    m, cl = c.facts.elements["multispeciality"], c.facts.elements["clinic_framing"]
    if m and cl:
        return "Yes, but also called it clinic/OPD"
    return "Yes" if m else ("No - called it clinic/OPD" if cl else "No")


def _hubs_value(c: CallFacts) -> str:
    if c.facts.hubs:
        return ", ".join(HUBS[h] for h in c.facts.hubs)
    return GENERIC_GURGAON if c.facts.generic_gurgaon else "None"


def master_table(calls: list[CallFacts]) -> pd.DataFrame:
    rows = []
    for c in calls:
        f, m = c.facts, c.meta
        labels = [DOUBTS[k][0] for k in DOUBTS if k in f.doubts]
        cust = "\n".join(f"[{DOUBTS[k][0]}] {f.doubts[k][0].customer}" for k in DOUBTS if k in f.doubts)
        if f.doubts:
            agent = "\n".join(f"[{DOUBTS[k][0]}] {f.doubts[k][0].agent or '(no agent reply captured)'}"
                              for k in DOUBTS if k in f.doubts)
        else:
            agent = f.pitch_line
        rows.append({
            "Call ID": c.call_id,
            "Call Date": _date_str(m),
            "Talktime (mins)": _minutes(m.get("talktime")),
            "Agent ID": m.get("agent_id") or m.get("agent_name") or "",
            "Customer City": m.get("city") or "",
            "Pitched? (Y/N)": "Y" if f.pitched else "N",
            "Hubs Mentioned": _hubs_value(c),
            "Multispeciality?": _multi_value(c),
            "No Overcharge Pitch?": _yn(c, "ethical"),
            "App Pushed?": _yn(c, "app"),
            "Customer Doubt Category": "; ".join(labels) or "None",
            "Customer Quote (Verbatim)": cust,
            "Agent Quote (Verbatim)": agent,
            # audit columns
            "Zero Wait Pitch?": _yn(c, "zero_wait"),
            "Family Discount?": _yn(c, "family_discount"),
            "Cashless?": _yn(c, "cashless"),
            "Hospital Raised By": "Agent" if f.pitched else ("Customer only" if f.any_mention else "Not discussed"),
            "Search Terms Matched": ", ".join(m.get("matched_terms") or []),
        })
    df = pd.DataFrame(rows)
    if not df.empty:
        df = df.sort_values(["Call Date", "Call ID"], ascending=[False, True], kind="stable").reset_index(drop=True)
    return df


def _pick_examples(items: list[tuple[CallFacts, object]], n: int) -> list[tuple[CallFacts, object]]:
    """Most recent first, one per call, readable length preferred."""
    seen, out = set(), []
    ranked = sorted(items, key=lambda ce: _date_str(ce[0].meta), reverse=True)
    ranked = sorted(ranked, key=lambda ce: not (15 <= len(ce[1].customer) <= 300))  # stable: keeps recency order
    for c, ev in ranked:
        if c.call_id in seen:
            continue
        seen.add(c.call_id)
        out.append((c, ev))
        if len(out) == n:
            break
    return out


def doubt_tables(calls: list[CallFacts]) -> tuple[pd.DataFrame, pd.DataFrame]:
    pitched = [c for c in calls if c.facts.pitched]
    n_p, n_all = len(pitched), len(calls)
    scripts = _counter_scripts()
    summary, detail = [], []
    for key, (label, meaning) in DOUBTS.items():
        hit = [c for c in pitched if key in c.facts.doubts]
        evid = [(c, ev) for c in hit for ev in c.facts.doubts[key]]
        quotes = _pick_examples(evid, 5)
        answers = _pick_examples([(c, ev) for c, ev in evid if ev.agent], 2)
        summary.append({"Customer Doubt": label, "What It Sounds Like": meaning, "Pitched Calls": len(hit),
                        "% of Pitched Calls": pct(len(hit), n_p), "% of All Analyzed Calls": pct(len(hit), n_all)})
        detail.append({
            "Doubt Category": label,
            "What Customer Actually Asks (Verbatim Quotes)": "\n\n".join(
                f'"{ev.customer}"  (Call {c.call_id}, {_date_str(c.meta)[:10]})' for c, ev in quotes)
            or "No occurrences in this window.",
            "Frequency (Call Count)": len(hit),
            "% of Pitched Calls": pct(len(hit), n_p),
            "How Agents Currently Answer (Real Example)": "\n\n".join(
                f'Customer: "{ev.customer}"\nAgent: "{ev.agent}"\n(Call {c.call_id})' for c, ev in answers)
            or ("No agent reply captured." if hit else "No occurrences in this window."),
            "Recommended Counter-Script for Sales Deck": scripts[key],
        })
    any_doubt = sum(1 for c in pitched if c.facts.doubts)
    s = pd.DataFrame(summary).sort_values("Pitched Calls", ascending=False, kind="stable")
    s = pd.concat([s, pd.DataFrame([{"Customer Doubt": "Any doubt (at least one of the above)", "What It Sounds Like": "",
                                     "Pitched Calls": any_doubt, "% of Pitched Calls": pct(any_doubt, n_p),
                                     "% of All Analyzed Calls": pct(any_doubt, n_all)}])], ignore_index=True)
    d = pd.DataFrame(detail).sort_values("Frequency (Call Count)", ascending=False, kind="stable").reset_index(drop=True)
    return s, d


def element_table(calls: list[CallFacts]) -> pd.DataFrame:
    pitched = [c for c in calls if c.facts.pitched]
    n_p = len(pitched)

    def row(label: str, said: list[CallFacts], quote_of, negative: bool = False) -> dict:
        ex = next(((c, quote_of(c)) for c in sorted(said, key=lambda c: _date_str(c.meta), reverse=True) if quote_of(c)), None)
        return {"Pitch Element": label, "Pitched Calls Where Said": len(said), "% of Pitched Calls": pct(len(said), n_p),
                "Pitched Calls Where Missing": None if negative else n_p - len(said),
                "Real Example (Agent Verbatim)": ex[1] if ex else "", "Example Call ID": ex[0].call_id if ex else ""}

    el = lambda k: [c for c in pitched if c.facts.elements.get(k)]  # noqa: E731
    rows = [row(ELEMENTS["multispeciality"], el("multispeciality"), lambda c: c.facts.element_quotes.get("multispeciality")),
            row(ELEMENTS["clinic_framing"], el("clinic_framing"), lambda c: c.facts.element_quotes.get("clinic_framing"), True)]
    for h, label in HUBS.items():
        rows.append(row(f"Hub named: {label}", [c for c in pitched if h in c.facts.hubs], lambda c, h=h: c.facts.hub_quotes.get(h)))
    rows.append(row("Any of the 5 hubs named", [c for c in pitched if c.facts.hubs],
                    lambda c: next(iter(c.facts.hub_quotes.values()), "")))
    rows.append(row('Only generic "Gurgaon" (no hub named)', [c for c in pitched if c.facts.generic_gurgaon], lambda c: "", True))
    for k in ("ethical", "zero_wait", "app", "family_discount", "cashless"):
        rows.append(row(ELEMENTS[k], el(k), lambda c, k=k: c.facts.element_quotes.get(k)))
    return pd.DataFrame(rows)


def hub_table(calls: list[CallFacts]) -> pd.DataFrame:
    pitched = [c for c in calls if c.facts.pitched]
    n_p = len(pitched)
    counts = {h: sum(h in c.facts.hubs for c in pitched) for h in HUBS}
    total_mentions = sum(counts.values())
    rows = [{"Hub": label, "Pitched Calls Where Agent Named It": counts[h], "% of Pitched Calls": pct(counts[h], n_p),
             "Share of All Hub Mentions": pct(counts[h], total_mentions)} for h, label in HUBS.items()]
    generic = sum(c.facts.generic_gurgaon for c in pitched)
    none = sum(not c.facts.hubs and not c.facts.generic_gurgaon for c in pitched)
    rows.append({"Hub": 'Only generic "Gurgaon" (no hub)', "Pitched Calls Where Agent Named It": generic,
                 "% of Pitched Calls": pct(generic, n_p), "Share of All Hub Mentions": None})
    rows.append({"Hub": "No location named at all", "Pitched Calls Where Agent Named It": none,
                 "% of Pitched Calls": pct(none, n_p), "Share of All Hub Mentions": None})
    return pd.DataFrame(rows)


def gap_table(elements: pd.DataFrame) -> pd.DataFrame:
    """The 3 core-story items said in the fewest pitched calls."""
    core = [ELEMENTS[k] for k in ("multispeciality", "ethical", "zero_wait", "app", "family_discount", "cashless")]
    core.append("Any of the 5 hubs named")
    g = elements[elements["Pitch Element"].isin(core)].sort_values(
        ["Pitched Calls Where Said", "Pitch Element"], kind="stable").head(3).reset_index(drop=True)
    g["Pitch Element"] = [f"{i}. {name}" for i, name in enumerate(g["Pitch Element"], 1)]
    g = g.rename(columns={"Pitch Element": "What Agents Are Missing"})
    return g[["What Agents Are Missing", "Pitched Calls Where Said", "% of Pitched Calls", "Pitched Calls Where Missing"]]


def headline(calls: list[CallFacts], run: RunInfo) -> list[tuple[str, object, str]]:
    n = len(calls)
    pitched = [c for c in calls if c.facts.pitched]
    p = len(pitched)
    rows: list[tuple[str, object, str]] = [
        ("Window", f"{run.start} to {run.end}", "text"),
        ("Data source", run.source, "text"),
        ("Calls scanned (unique calls returned by search)", run.scanned, "int"),
        ("Calls analyzed (full transcript fetched)", n, "int"),
    ]
    if run.failed or run.empty:
        rows.append(("Calls skipped (transcript failed or empty)", run.failed + run.empty, "int"))
    rows += [
        ("Calls where the agent pitched PB Hospital", p, "int"),
        ("Hospital pitch rate (pitched / analyzed)", pct(p, n), "pct"),
    ]
    if run.qualified_ncr is not None:
        rows += [("All qualified NCR health sales calls in window (portal count)", run.qualified_ncr, "int"),
                 ("Pitched calls / all qualified NCR health sales calls", pct(p, run.qualified_ncr), "pct")]
    cust_only = sum(1 for c in calls if c.facts.any_mention and not c.facts.pitched)
    any_doubt = sum(1 for c in pitched if c.facts.doubts)
    clinic = sum(1 for c in pitched if c.facts.elements["clinic_framing"])
    rows += [
        ("Calls where only the customer brought up PB Hospital", cust_only, "int"),
        ("Pitched calls where the customer raised at least one doubt", any_doubt, "int"),
        ("  as % of pitched calls", pct(any_doubt, p), "pct"),
        ('Pitched calls where the agent called it a "clinic" / "OPD pass"', clinic, "int"),
    ]
    return rows


# --------------------------------------------------------------------------- #
# Workbook
# --------------------------------------------------------------------------- #
NAVY, TEAL, LIGHT, GREY = "0B2A4A", "00A3A1", "EEF6F8", "D0D7DE"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
BAND_FILL = PatternFill("solid", fgColor=LIGHT)
BORDER = Border(*(Side(style="thin", color=GREY),) * 4)
TITLE_FONT = Font(bold=True, size=16, color=NAVY)
SECTION_FONT = Font(bold=True, size=12, color=TEAL)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
PCT_FMT = "0.0%"


def _is_pct_col(name: str) -> bool:
    return name.startswith("%") or name.startswith("Share")


def _write_df(ws, df: pd.DataFrame, top: int, left: int = 1) -> int:
    """Write df with a styled header at row `top`. Returns the next free row."""
    for j, col in enumerate(df.columns):
        cell = ws.cell(row=top, column=left + j, value=col)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, rec in enumerate(df.itertuples(index=False), 1):
        for j, (col, val) in enumerate(zip(df.columns, rec)):
            if val is None or (isinstance(val, float) and math.isnan(val)):
                val = None
            cell = ws.cell(row=top + i, column=left + j, value=val)
            cell.border, cell.alignment = BORDER, WRAP_TOP
            if _is_pct_col(col) and isinstance(val, (int, float)):
                cell.number_format = PCT_FMT
            if i % 2 == 0:
                cell.fill = BAND_FILL
    return top + len(df) + 1


def _fit_rows(ws, first_row: int, last_row: int, widths: dict[int, float]) -> None:
    """Approximate row heights for wrapped text (openpyxl cannot auto-fit)."""
    for r in range(first_row, last_row + 1):
        lines = 1
        for col, w in widths.items():
            v = ws.cell(row=r, column=col).value
            if isinstance(v, str) and v:
                per_line = max(int(w * 1.15), 1)
                lines = max(lines, sum(max(1, math.ceil(len(part) / per_line)) for part in v.split("\n")))
        if lines > 1:
            ws.row_dimensions[r].height = min(15 * lines + 4, 409)


def _set_widths(ws, widths: list[float]) -> dict[int, float]:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return {i: w for i, w in enumerate(widths, 1)}


def build_workbook(calls: list[CallFacts], run: RunInfo, out: Path) -> dict:
    master = master_table(calls)
    doubt_summary, doubt_detail = doubt_tables(calls)
    elements = element_table(calls)
    hubs = hub_table(calls)
    gaps = gap_table(elements)
    head = headline(calls, run)

    wb = Workbook()
    # ---- Tab 1: Executive Summary ----
    ws = wb.active
    ws.title = "Executive Summary"
    widths = _set_widths(ws, [58, 36, 20, 20, 20])
    ws["A1"] = "PB Health Hospital - Sales Call Analysis"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Real calls, {run.start} to {run.end}. Generated {datetime.now():%Y-%m-%d %H:%M}."
    ws["A2"].font = Font(italic=True, color="5B6B79")
    r = 4
    ws.cell(row=r, column=1, value="Headline numbers").font = SECTION_FONT
    r += 1
    for j, h in enumerate(("Metric", "Value"), 1):
        c = ws.cell(row=r, column=j, value=h)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
    for metric, val, kind in head:
        r += 1
        ws.cell(row=r, column=1, value=metric).border = BORDER
        c = ws.cell(row=r, column=2, value=val)
        c.border, c.alignment = BORDER, Alignment(horizontal="left", wrap_text=True)
        if kind == "pct":
            c.number_format = PCT_FMT
            if val is None:
                c.value = "n/a (no pitched calls)"
        elif kind == "int":
            c.number_format = "#,##0"
    r += 2
    ws.cell(row=r, column=1, value="Top 3 pitch gaps (core story items said in the fewest pitched calls)").font = SECTION_FONT
    r = _write_df(ws, gaps, r + 1)
    five = hubs.iloc[:5]
    never = five[five["Pitched Calls Where Agent Named It"] == 0]["Hub"].tolist()
    if never:
        hub_line = "Hubs never named in a pitched call: " + ", ".join(never)
    else:
        least = five.sort_values("Pitched Calls Where Agent Named It", kind="stable").head(2)
        hub_line = "Hubs named least: " + "; ".join(
            f"{h} ({n} calls)" for h, n in zip(least["Hub"], least["Pitched Calls Where Agent Named It"]))
    ws.cell(row=r, column=1, value=hub_line)
    ws.cell(row=r, column=1).font = Font(italic=True)
    r += 2
    ws.cell(row=r, column=1, value="Hub location mentions (pitched calls)").font = SECTION_FONT
    r = _write_df(ws, hubs, r + 1) + 1
    ws.cell(row=r, column=1, value="Customer doubts (pitched calls)").font = SECTION_FONT
    r = _write_df(ws, doubt_summary, r + 1) + 1
    ws.cell(row=r, column=1, value="How these numbers are counted").font = SECTION_FONT
    notes = [
        "Pitched = an AGENT line names PB Health / PB Hospital / Policybazaar ka hospital / khud ka / humara / apna "
        "hospital (\"aap apna hospital choose kar sakte hain\" is excluded).",
        "Hubs, story elements and doubts are only counted inside the part of the call where PB Hospital is being "
        "discussed (1 line before each mention to 6 after, extended while the talk stays on the hospital). This stops "
        "policy talk about cashless / premium discount / policy app from being counted.",
        "Doubts come from CUSTOMER lines only; story elements and hubs from AGENT lines only. One call can have "
        "several doubts. Every count can be checked line by line in the 'Call-by-Call Master Data' tab.",
        *run.notes,
    ]
    for note in notes:
        r += 1
        c = ws.cell(row=r, column=1, value=note)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        ws.row_dimensions[r].height = 15 * max(1, math.ceil(len(note) / 150)) + 4
    _fit_rows(ws, 5, r - len(notes) - 1, {k: v for k, v in widths.items() if k > 1})

    # ---- Tab 2: Customer Doubts ----
    ws = wb.create_sheet("Customer Doubts & Objections")
    widths = _set_widths(ws, [24, 70, 13, 12, 70, 70])
    end = _write_df(ws, doubt_detail, 1)
    ws.freeze_panes = "B2"
    _fit_rows(ws, 2, end - 1, widths)

    # ---- Tab 3: Pitch Elements ----
    ws = wb.create_sheet("Pitch Elements Tracker")
    widths = _set_widths(ws, [46, 16, 14, 16, 80, 16])
    ws["A1"] = f"Pitch elements, counted over {sum(c.facts.pitched for c in calls)} pitched calls"
    ws["A1"].font = SECTION_FONT
    end = _write_df(ws, elements, 3)
    ws.freeze_panes = "B4"
    _fit_rows(ws, 4, end - 1, widths)

    # ---- Tab 4: Master Data ----
    ws = wb.create_sheet("Call-by-Call Master Data")
    widths = _set_widths(ws, [14, 17, 10, 14, 16, 9, 30, 22, 12, 10, 34, 70, 70, 10, 10, 10, 14, 40])
    end = _write_df(ws, master, 1)
    ws.freeze_panes = "B2"
    if not master.empty:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(master.columns))}{end - 1}"
    for row in ws.iter_rows(min_row=2, max_row=end - 1):
        for cell in row:
            if not isinstance(cell.value, str) or len(cell.value) < 40:
                cell.alignment = Alignment(vertical="top")
    _fit_rows(ws, 2, end - 1, {12: widths[12], 13: widths[13]})

    for sheet in wb.worksheets:   # print: landscape, fit all columns on one page width
        sheet.page_setup.orientation = "landscape"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {"master": master, "doubt_summary": doubt_summary, "doubt_detail": doubt_detail,
            "elements": elements, "hubs": hubs, "gaps": gaps, "headline": head}


def write_dialogue_csv(calls: list[CallFacts], out: Path) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["call_id", "call_date", "transcript", "line", "speaker", "in_hospital_window",
                    "text_verbatim", "text_harmonized"])
        for c in calls:
            inside = hospital_window(c.turns)
            for t, in_w in zip(c.turns, inside):
                w.writerow([c.call_id, _date_str(c.meta), c.source_transcript, t.idx + 1, t.speaker,
                            "Y" if in_w else "", t.text, t.norm])


# --------------------------------------------------------------------------- #
# Console
# --------------------------------------------------------------------------- #
def _fmt_val(val: object, kind: str) -> str:
    if kind == "pct":
        return "n/a" if val is None else f"{val:.1%}"
    if kind == "int":
        return f"{val:,}"
    return str(val)


def print_summary(tables: dict, out: Path) -> None:
    line = "=" * 78
    print(f"\n{line}\nPB HEALTH HOSPITAL - SALES CALL ANALYSIS\n{line}")
    for metric, val, kind in tables["headline"]:
        print(f"  {metric:<64} {_fmt_val(val, kind)}")
    print(f"\nTOP CUSTOMER DOUBTS (pitched calls)\n{'-' * 78}")
    print(f"  {'Doubt':<40} {'Calls':>7} {'% pitched':>11} {'% all':>9}")
    for rec in tables["doubt_summary"].itertuples(index=False):
        p = "n/a" if rec[3] is None or pd.isna(rec[3]) else f"{rec[3]:.1%}"
        a = "n/a" if rec[4] is None or pd.isna(rec[4]) else f"{rec[4]:.1%}"
        print(f"  {rec[0][:40]:<40} {rec[2]:>7,} {p:>11} {a:>9}")
    print(f"\nHUB MENTIONS (pitched calls)\n{'-' * 78}")
    for rec in tables["hubs"].itertuples(index=False):
        p = "n/a" if rec[2] is None or pd.isna(rec[2]) else f"{rec[2]:.1%}"
        print(f"  {rec[0][:40]:<40} {rec[1]:>7,} {p:>11}")
    print(f"\nTOP 3 PITCH GAPS\n{'-' * 78}")
    for rec in tables["gaps"].itertuples(index=False):
        p = "n/a" if rec[2] is None or pd.isna(rec[2]) else f"{rec[2]:.1%}"
        print(f"  {rec[0][:53]:<53} said in {rec[1]:>5,} calls ({p})")
    print(f"\nExcel saved to: {out.resolve()}\n{line}\n")


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def run(args: argparse.Namespace) -> int:
    today = date.today()
    start = date.fromisoformat(args.date_from) if args.date_from else today - timedelta(days=args.days)
    end = date.fromisoformat(args.date_to) if args.date_to else today
    if start > end:
        print(f"--from {start} is after --to {end}", file=sys.stderr)
        return 2

    if args.from_db:
        if not Path(args.from_db).exists():
            print(f"No database at {args.from_db}", file=sys.stderr)
            return 2
        records, info = load_from_db(Path(args.from_db), start, end)
    elif args.cached:
        records, info = load_from_cache(start, end)
    else:
        import requests

        from api_client import PortalError
        try:
            records, info = load_from_live(start, end, args.max_calls, not args.no_category,
                                           not args.no_denominator, args.workers)
        except requests.exceptions.ConnectionError as exc:
            print("Cannot reach the PB portal API (datascienceapi.policybazaar.com). It only resolves on the "
                  f"PB office network / VPN - connect and re-run.\n  detail: {exc}", file=sys.stderr)
            return 2
        except PortalError as exc:
            print(f"Portal error: {exc}", file=sys.stderr)
            return 2

    calls = [analyze_call(r.call_id, r.english, r.hindi, r.meta) for r in records]
    calls = [c for c in calls if c.turns]
    info.empty = len(records) - len(calls)
    if not calls:
        print(f"No calls with transcripts found for {start} to {end} ({info.source}). Nothing to export.",
              file=sys.stderr)
        return 1
    out = Path(args.out)
    tables = build_workbook(calls, info, out)
    write_dialogue_csv(calls, out.with_name(OUT_CSV.name))
    print_summary(tables, out)
    print(f"Harmonized dialogue (every line, verbatim + normalized): {out.with_name(OUT_CSV.name).resolve()}")
    return 0


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Analyse real PB Health Hospital sales calls and export to Excel.")
    ap.add_argument("--days", type=int, default=21, help="look back N days up to today (default 21)")
    ap.add_argument("--from", dest="date_from", help="start date YYYY-MM-DD (overrides --days)")
    ap.add_argument("--to", dest="date_to", help="end date YYYY-MM-DD (default today)")
    ap.add_argument("--max-calls", type=int, help="analyse at most N calls (most recent first)")
    ap.add_argument("--cached", action="store_true", help="no API: use transcripts already in data/transcripts/")
    ap.add_argument("--from-db", nargs="?", const=str(config.DB_PATH), metavar="PATH",
                    help="no API: use the analyzer's SQLite DB (default data/pb_health.db), demo rows excluded")
    ap.add_argument("--no-category", action="store_true", help=f"skip the category-only ({CATEGORY_LABEL}) search")
    ap.add_argument("--no-denominator", action="store_true", help="skip the daily qualified-NCR-call count queries")
    ap.add_argument("--workers", type=int, default=config.TRANSCRIPT_WORKERS, help="parallel transcript fetches")
    ap.add_argument("--out", default=str(OUT_XLSX), help=f"output workbook (default {OUT_XLSX})")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO,
                        format="%(asctime)s %(levelname)s %(name)s: %(message)s")
    return run(args)


if __name__ == "__main__":
    sys.exit(main())
