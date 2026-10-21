"""
Build output/PB_Health_Call_Analysis.xlsx from the cached real calls.

Tab 4 (Call-by-Call Master Data) holds one row per call. Every number on tabs
1-3 is a COUNTIFS formula over Tab 4, so any figure can be traced back to the
calls behind it with an AutoFilter. The console summary evaluates the very
same criteria in pandas, so the two always agree.
"""
from __future__ import annotations

import hashlib
import json
import sqlite3
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import call_facts as cf
import config
from deck_generator import DECK

SUMMARY, DOUBTS, ELEMENTS, MASTER = ("Executive Summary", "Customer Doubts & Objections",
                                     "Pitch Elements Tracker", "Call-by-Call Master Data")
M = f"'{MASTER}'"

MASTER_COLS = [
    # the 13 columns asked for, in order
    "Call ID", "Call Date", "Talktime (mins)", "Agent ID", "Customer City", "Pitched? (Y/N)", "Hubs Mentioned",
    "Multispeciality?", "No Overcharge Pitch?", "App Pushed?", "Customer Doubt Category",
    "Customer Quote (Verbatim)", "Agent Quote (Verbatim)",
    # audit columns the summary formulas count on
    "Named One of 5 Hubs?", "Zero Wait Pitch?", "Family Discount?", "Cashless Pitch?", "Other Agent Claims",
    "Customer in Delhi NCR?", "Hospital Raised By", "Pitch Phrase Used", "Said 'PB Health' (any context)?", "Search Cluster",
    "Transcript Analysed?", "Customer State", "Lead ID",
]
MS_YES, MS_OPD, MS_NONE = "Yes", "No - only OPD/consultation/clinic words", "No - scope not described"
# Lead city -> inside the area the hubs serve (config.NCR_CITIES + NCR satellite towns).
NCR_CITIES = {c.lower() for c in config.NCR_CITIES} | {
    "sonipat", "bahadurgarh", "palwal", "ballabgarh", "ballabhgarh", "greater noida west", "manesar", "jhajjar"}

NAVY, ORANGE, LIGHT, GREY = "0B2A4A", "F26B21", "EEF6F8", "5B6B79"
FONT = "Arial"
THIN = Side(style="thin", color="D5DDE3")


# --------------------------------------------------------------------------- #
# Data
# --------------------------------------------------------------------------- #
def _yn(flag: bool) -> str:
    return "Y" if flag else "N"


def _clean(v):
    """Speech-to-text occasionally emits control characters that Excel files cannot hold."""
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def load_rows(conn: sqlite3.Connection, start: date, end: date) -> tuple[pd.DataFrame, dict[str, cf.CallFacts]]:
    q = """SELECT h.call_id, h.call_ts, h.agent_id, h.city, h.state, h.talktime, h.lead_id, h.terms,
                  t.hindi, t.english, t.fetched_at
           FROM hits h LEFT JOIN transcripts t ON t.call_id = h.call_id"""
    rows, facts = [], {}
    for r in conn.execute(q):
        day = (r["call_ts"] or "")[:10]
        if not day or not (str(start) <= day <= str(end)):
            continue
        hindi = json.loads(r["hindi"] or "[]")
        english = json.loads(r["english"] or "[]")
        has_t = r["fetched_at"] is not None and bool(hindi or english)
        f = cf.analyse_call(r["call_id"], hindi, english) if has_t else None
        if f:
            facts[r["call_id"]] = f
        pitched = bool(f and f.pitched)
        dash = "-" if not pitched else None
        doubts = [cf.DOUBT_LABELS[d] for d in f.doubts] if f else []
        first = f.doubts[0] if f and f.doubts else None
        if f and f.pitched:
            ms = MS_YES if f.elements["multispeciality"] else MS_OPD if f.elements["clinic_opd"] else MS_NONE
        rows.append({
            "Call ID": int(r["call_id"]) if str(r["call_id"]).isdigit() else r["call_id"],
            "Call Date": datetime.strptime(day, "%Y-%m-%d").date(),
            "Talktime (mins)": round((r["talktime"] or 0) / 60, 1),
            "Agent ID": r["agent_id"],
            "Customer City": r["city"],
            "Pitched? (Y/N)": _yn(pitched) if f else "n/a",
            "Hubs Mentioned": dash or ("; ".join(f.locations) or "None"),
            "Multispeciality?": dash or ms,
            "No Overcharge Pitch?": dash or _yn(f.elements["ethical"]),
            "App Pushed?": _yn(f.elements["app"]) if f else "n/a",
            "Customer Doubt Category": ("; ".join(doubts) or "None") if f and f.hospital_mentioned else "-",
            "Customer Quote (Verbatim)": f.doubt_quotes[first] if first else "",
            "Agent Quote (Verbatim)": (f"[Reply] {f.doubt_replies[first]}" if first and f.doubt_replies[first]
                                       else f"[Pitch] {f.pitch_quote}" if f and f.pitch_quote else ""),
            "Named One of 5 Hubs?": dash or _yn(any(h in f.locations for h in cf.FIVE_HUBS)),
            "Zero Wait Pitch?": dash or _yn(f.elements["zero_wait"]),
            "Family Discount?": dash or _yn(f.elements["family_discount"]),
            "Cashless Pitch?": dash or _yn(f.elements["cashless"]),
            "Other Agent Claims": dash or ("; ".join(f.claims) or "None"),
            "Customer in Delhi NCR?": ("Y" if r["city"].strip().lower() in NCR_CITIES else "N") if r["city"] else "Unknown",
            "Hospital Raised By": f.mentioned_by if f else "-",
            "Pitch Phrase Used": "; ".join(f.pitch_phrases) if f and f.pitch_phrases else "-",
            "Said 'PB Health' (any context)?": _yn(f.pb_health_said) if f else "n/a",
            "Search Cluster": "; ".join(sorted(json.loads(r["terms"]))),
            "Transcript Analysed?": _yn(has_t),
            "Customer State": r["state"],
            "Lead ID": r["lead_id"],
        })
    df = pd.DataFrame(rows, columns=MASTER_COLS).map(_clean)
    order = {"Y": 0, "N": 1, "n/a": 2}
    df["_o1"] = df["Pitched? (Y/N)"].map(order)
    df["_o2"] = df["Hospital Raised By"].ne("-").map({True: 0, False: 1})
    df = df.sort_values(["_o1", "_o2", "Call Date", "Call ID"], ascending=[True, True, False, True])
    return df.drop(columns=["_o1", "_o2"]).reset_index(drop=True), facts


# --------------------------------------------------------------------------- #
# COUNTIFS that exist twice: as an Excel formula and as a pandas value
# --------------------------------------------------------------------------- #
@dataclass
class Counter:
    df: pd.DataFrame

    def rng(self, col: str) -> str:
        c = get_column_letter(MASTER_COLS.index(col) + 1)
        return f"{M}!${c}$2:${c}${len(self.df) + 1}"

    def _mask(self, col: str, crit: str) -> pd.Series:
        s = self.df[col].astype(str)
        if crit.startswith("*") and crit.endswith("*"):
            return s.str.lower().str.contains(crit[1:-1].lower(), regex=False)
        if crit.startswith("<>"):
            return s.str.lower() != crit[2:].lower()
        return s.str.lower() == crit.lower()

    def count(self, *conds: tuple[str, str]) -> tuple[str, int]:
        formula = "COUNTIFS(" + ",".join(f'{self.rng(c)},"{v}"' for c, v in conds) + ")"
        mask = pd.Series(True, index=self.df.index)
        for c, v in conds:
            mask &= self._mask(c, v)
        return formula, int(mask.sum())


PITCHED = ("Pitched? (Y/N)", "Y")
PBH_COL = "Said 'PB Health' (any context)?"


# --------------------------------------------------------------------------- #
# Styling helpers
# --------------------------------------------------------------------------- #
def _font(bold=False, size=10, color="1E2A35", italic=False) -> Font:
    return Font(name=FONT, bold=bold, size=size, color=color, italic=italic)


def _header(ws, row: int, labels: list[str], col: int = 1) -> None:
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=label)
        c.font = _font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=THIN)


def _section(ws, row: int, text: str) -> None:
    c = ws.cell(row=row, column=1, value=text)
    c.font = _font(bold=True, size=11, color=ORANGE)


def _body(cell, wrap=True, fmt: str | None = None, bold=False) -> None:
    cell.font = _font(bold=bold)
    cell.alignment = Alignment(wrap_text=wrap, vertical="top")
    cell.border = Border(bottom=THIN)
    if fmt:
        cell.number_format = fmt


def _fit_rows(ws, first: int, last: int, widths: dict[int, float]) -> None:
    """Rough row heights for wrapped text (openpyxl cannot autofit)."""
    for r in range(first, last + 1):
        lines = 1
        for col, w in widths.items():
            v = ws.cell(row=r, column=col).value
            if isinstance(v, str) and not v.startswith("="):
                est = sum(max(1, -(-len(part) // max(int(w * 1.1), 1))) for part in v.split("\n"))
                lines = max(lines, est)
        ws.row_dimensions[r].height = min(409, 14 * lines + 4)


# --------------------------------------------------------------------------- #
# Workbook
# --------------------------------------------------------------------------- #
def _pick_examples(items: list[tuple[str, str, str, str]], k: int) -> list[tuple[str, str, str, str]]:
    """Deterministic spread of examples (hash order) - not cherry-picked by length or tone."""
    return sorted(items, key=lambda x: hashlib.md5(x[0].encode()).hexdigest())[:k]


COUNTER_SCRIPTS = {  # deck_generator.DECK['objections'] index + ops caveat
    "scope": (3, ""),
    "trust": (0, ""),
    "location": (1, "CONFIRM WITH OPS FIRST: on these calls agents say only the MG Road, Gurgaon hospital is open and "
                    "the rest are 'coming soon'. Only name hubs that are open today."),
    "cashless": (2, "CONFIRM WITH OPS: the list of insurers with cashless tie-ups."),
    "discount": (4, "CONFIRM WITH OPS: exact redemption steps. Agents quote 50%-75% off on calls - one approved number is needed."),
    "skeptic": (6, "Say it once in one line and move on; don't repeat the pitch after a 'no'."),
}


def build(conn: sqlite3.Connection, start: date, end: date, out: Path) -> int:
    df, facts = load_rows(conn, start, end)
    if df.empty:
        print("No cached calls in the date range - run without --skip-fetch first.")
        return 1
    k = Counter(df)
    period = f"{df['Call Date'].min()} to {df['Call Date'].max()}"
    gaps = [tuple(r) for r in conn.execute(
        "SELECT cluster, w_from, w_to FROM gaps WHERE substr(w_from, 1, 10) BETWEEN ? AND ? ORDER BY w_from",
        (str(start), str(end)))]

    out.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name=MASTER, index=False)
        wb = xw.book
        ws_sum = wb.create_sheet(SUMMARY, 0)
        ws_dbt = wb.create_sheet(DOUBTS, 1)
        ws_el = wb.create_sheet(ELEMENTS, 2)
        console = _write_elements(ws_el, k)
        doubt_rows = _write_doubts(ws_dbt, k, df, facts)
        _write_summary(ws_sum, k, df, period, console, doubt_rows, gaps)
        _style_master(wb[MASTER], df)
        wb.calculation.fullCalcOnLoad = True
        wb.active = 0

    _print_console(k, df, period, doubt_rows, out, gaps)
    return 0


# ---- Tab 3 -----------------------------------------------------------------
ELEMENT_ROWS = [
    # (section, label, conditions, rule shown to the reader)
    ("CORE STORY", "Real Multispeciality Hospital", [("Multispeciality?", MS_YES)],
     "Agent used hospital-scope words near the pitch: multispeciality, ICU, OT, surgery/admission, IPD, day care."),
    ("CORE STORY", "Described only as OPD / consultation / clinic", [("Multispeciality?", MS_OPD)],
     "Agent talked about the hospital only as OPD / consultation / clinic - no hospital-scope words."),
    ("CORE STORY", "Hospital scope not described at all", [("Multispeciality?", MS_NONE)],
     "Agent named the hospital but said nothing about what it treats."),
    ("LOCATIONS", "IFFCO Chowk", [("Hubs Mentioned", "*IFFCO Chowk*")], "Agent said 'IFFCO' near the pitch."),
    ("LOCATIONS", "Sector 62 / Golf Course", [("Hubs Mentioned", "*Sector 62 / Golf Course*")], "'Sector 62' or 'Golf Course'."),
    ("LOCATIONS", "Faridabad Sec 2", [("Hubs Mentioned", "*Faridabad Sec 2*")], "'Faridabad' near the pitch."),
    ("LOCATIONS", "Noida Sec 50", [("Hubs Mentioned", "*Noida Sec 50*")], "'Noida' or 'Sector 50' near the pitch (includes 'Noida coming soon')."),
    ("LOCATIONS", "Rohini", [("Hubs Mentioned", "*Rohini*")], "'Rohini' near the pitch."),
    ("LOCATIONS", "Any of the 5 hubs by name", [("Named One of 5 Hubs?", "Y")], "At least one of the five hub names above."),
    ("LOCATIONS", "MG Road, Gurgaon (address only)", [("Hubs Mentioned", "*MG Road, Gurgaon*")],
     "Agent gave the MG Road / MG Road metro address instead of a hub name."),
    ("LOCATIONS", "Gurgaon (generic)", [("Hubs Mentioned", "*Gurgaon (generic)*")], "'Gurgaon' with no hub or road."),
    ("LOCATIONS", "No location named", [("Hubs Mentioned", "None")], "Hospital pitched with no place at all."),
    ("USPs", "Ethical Care / No Incentive to Overcharge", [("No Overcharge Pitch?", "Y")],
     "salary-based doctors, no commission/target, no overcharging, no unnecessary tests."),
    ("USPs", "Zero Waiting Time / Digital First", [("Zero Wait Pitch?", "Y")],
     "no waiting / no queue / pre-book / online appointment / digital records ('waiting period' excluded)."),
    ("USPs", "PB Health App Download", [("App Pushed?", "Y")],
     "Agent named the 'PB Health app' anywhere in the call, or told the customer to download/install an app near the pitch."),
    ("USPs", "Exclusive Family Discount", [("Family Discount?", "Y")], "discount/percent within ~80 characters of family/parents/wife/children."),
    ("USPs", "Cashless Guarantee", [("Cashless Pitch?", "Y")], "cashless / zero deduction / no deposit said about the hospital."),
]
CLAIM_ROWS = [(label, [("Other Agent Claims", f"*{tag}*")]) for label, tag in cf.CLAIM_TAGS.items()]
PHRASE_ROWS = ["PB Health Hospital / PB Hospital", "Policybazaar ka khud ka/apna hospital", "khud ka hospital",
               "Policybazaar hospital", "apna/hamara hospital", "PB Health ... hospital"]


def _write_elements(ws, k: Counter) -> dict:
    ws["A1"] = "Pitch Elements Tracker"
    ws["A1"].font = _font(bold=True, size=14, color=NAVY)
    ws["A2"] = "Base = calls where the agent pitched PB Hospital. Each count is a COUNTIFS over the Call-by-Call tab."
    ws["A2"].font = _font(italic=True, color=GREY)
    pitched_f, pitched_n = k.count(PITCHED)
    ws["A3"], ws["B3"] = "Pitched calls (base)", f"={pitched_f}"
    ws["A3"].font = _font(bold=True)
    ws["B3"].font = _font(bold=True)
    base = "$B$3"
    vals: dict[str, tuple[int, int]] = {}   # label -> (count, sheet row)
    row = 5
    section = None
    _header(ws, row, ["Pitch Element", "Calls (of pitched)", "% of Pitched Calls", "What counts (rule applied to the agent's words)"])
    row += 1
    for sec, label, conds, rule in ELEMENT_ROWS:
        if sec != section:
            section = sec
            _section(ws, row, sec)
            row += 1
        f, n = k.count(PITCHED, *conds)
        _put_count_row(ws, row, label, f, base, rule)
        vals[label] = (n, row)
        row += 1
    row += 1
    _section(ws, row, "WHAT ELSE AGENTS SAY ABOUT THE HOSPITAL (verbatim claims, counted)")
    row += 1
    for label, conds in CLAIM_ROWS:
        f, n = k.count(PITCHED, *conds)
        _put_count_row(ws, row, label, f, base, "Said by the agent within the hospital part of the call.")
        vals[label] = (n, row)
        row += 1
    row += 1
    _section(ws, row, "HOW AGENTS NAME THE HOSPITAL (pitch phrase tags; a call can use several)")
    row += 1
    for tag in PHRASE_ROWS:
        f, n = k.count(PITCHED, ("Pitch Phrase Used", f"*{tag}*"))
        _put_count_row(ws, row, tag, f, base, "Normalised from Hinglish variants (khud ke/ki, policy bazar, p b, hospitel ...).")
        vals[tag] = (n, row)
        row += 1
    for col, w in {1: 46, 2: 16, 3: 16, 4: 90}.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A6"
    return {"pitched": pitched_n, "rows": vals}


def _put_count_row(ws, row: int, label: str, formula: str, base: str, rule: str) -> None:
    ws.cell(row=row, column=1, value=label)
    ws.cell(row=row, column=2, value=f"={formula}")
    ws.cell(row=row, column=3, value=f"=IFERROR(B{row}/{base},0)")
    ws.cell(row=row, column=4, value=rule)
    for col in range(1, 5):
        _body(ws.cell(row=row, column=col), fmt="0.0%" if col == 3 else "#,##0" if col == 2 else None)


# ---- Tab 2 -----------------------------------------------------------------
def _write_doubts(ws, k: Counter, df: pd.DataFrame, facts: dict[str, cf.CallFacts]) -> list[dict]:
    cols = ["Doubt Category", "What Customer Actually Asks (Verbatim Quotes)", "Frequency (Call Count)",
            "% of Pitched Calls", "How Agents Currently Answer (Real Example)", "Recommended Counter-Script for Sales Deck"]
    _header(ws, 1, cols)
    pitched_ids = set(df.loc[df["Pitched? (Y/N)"] == "Y", "Call ID"].astype(str))
    dates = dict(zip(df["Call ID"].astype(str), df["Call Date"].astype(str)))
    rows = []
    for key, label in cf.DOUBT_LABELS.items():
        f, n = k.count(PITCHED, ("Customer Doubt Category", f"*{label}*"))
        items = [(cid, fx.doubt_quotes[key], fx.doubt_replies.get(key, ""), dates.get(cid, ""))
                 for cid, fx in facts.items() if cid in pitched_ids and key in fx.doubts]
        rows.append({"key": key, "label": label, "formula": f, "n": n, "items": items})
    rows.sort(key=lambda r: (-r["n"], list(cf.DOUBT_LABELS).index(r["key"])))
    pitched_all = k.rng("Pitched? (Y/N)")
    for i, r in enumerate(rows, start=2):
        quotes = _pick_examples(r["items"], 4)
        answers = _pick_examples([x for x in r["items"] if len(x[2].split()) >= 8], 2)
        idx, caveat = COUNTER_SCRIPTS[r["key"]]
        title, cust_says, you_say, tip = DECK["objections"][idx]
        script = you_say.replace("\\\"", "\"").strip('"') + f"\n\nCoaching note: {tip}" + (f"\n\n{caveat}" if caveat else "")
        r["row"] = i
        values = [
            f"{r['label']}\n“{cf.DOUBT_DESC[r['key']]}”",
            "\n\n".join(f"“{q}”\n— Call {cid} ({d})" for cid, q, _, d in quotes) or "No customer raised this on the analysed calls.",
            f"={r['formula']}",
            f'=IFERROR(C{i}/COUNTIF({pitched_all},"Y"),0)',
            "\n\n".join(f"“{a}”\n— Call {cid}" for cid, _, a, _ in answers) or "No agent reply captured (no customer raised this, or the agent moved on).",
            script,
        ]
        for c, v in enumerate(values, start=1):
            _body(ws.cell(row=i, column=c, value=_clean(v)), fmt="0.0%" if c == 4 else "#,##0" if c == 3 else None,
                  bold=(c == 1))
    widths = {1: 30, 2: 70, 3: 14, 4: 13, 5: 60, 6: 60}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    _fit_rows(ws, 2, len(rows) + 1, {2: 70, 5: 60, 6: 60})
    ws.freeze_panes = "B2"
    n_last = len(rows) + 3
    ws.cell(row=n_last, column=1, value=(
        "Quotes are verbatim from the portal's romanised Hinglish transcript (speech-to-text, so some words are "
        "mis-heard). Examples are a fixed spread across calls, not hand-picked. Counter-scripts come from "
        "deck_generator.py (the sales deck) so both stay the same."))
    ws.cell(row=n_last, column=1).font = _font(italic=True, color=GREY)
    return rows


# ---- Tab 1 -----------------------------------------------------------------
def _write_summary(ws, k: Counter, df: pd.DataFrame, period: str, el: dict, doubt_rows: list[dict],
                   gaps: list[tuple[str, str, str]]) -> None:
    E = f"'{ELEMENTS}'"
    ws["A1"] = "PB Health Hospital - Real Sales Call Analysis"
    ws["A1"].font = _font(bold=True, size=16, color=NAVY)
    ws["A2"] = (f"Calls dated {period} | Source: Policybazaar Data Science portal (live search + transcript API) | "
                f"Health sales calls, category 10693, talktime > 2 min | Generated {datetime.now():%Y-%m-%d %H:%M}")
    ws["A2"].font = _font(italic=True, color=GREY)

    r = 4
    _section(ws, r, "HEADLINE COUNTS")
    r += 1
    _header(ws, r, ["Metric", "Value", "How it is counted"])
    r += 1
    pitched_cell = None
    analysed_cell = None
    headline = [
        ("Total calls scanned", f'=COUNTA({k.rng("Call ID")})', "#,##0",
         "Unique calls returned by the 15 search phrases (brand, Hinglish 'apna/khud ka hospital', hub names)."),
        ("Calls with a transcript analysed", f'={k.count(("Transcript Analysed?", "Y"))[0]}', "#,##0",
         "Full turn-by-turn transcript downloaded and read."),
        ("Calls where PB Hospital came up (agent or customer)",
         "=" + "+".join(k.count(("Hospital Raised By", v))[0] for v in ("Agent", "Customer", "Both")), "#,##0",
         "Policybazaar's own hospital named by either side."),
        ("Calls where the AGENT pitched PB Hospital", f"={k.count(PITCHED)[0]}", "#,##0",
         "Agent named Policybazaar's own hospital ('PB Health hospital', 'policybazaar ka khud ka hospital', ...)."),
        ("Hospital pitch rate (% of analysed calls)", None, "0.0%",
         "Pitched / analysed. The base is calls that matched the search phrases, not all sales calls."),
        ("Agent said 'PB Health' but never the hospital", "=" + k.count((PBH_COL, "Y"), ("Pitched? (Y/N)", "N"))[0],
         "#,##0", "'PB Health' used only for the app / wellness / membership - not counted as a hospital pitch."),
        ("Pitched calls naming one of the 5 hubs", f'={k.count(PITCHED, ("Named One of 5 Hubs?", "Y"))[0]}', "#,##0",
         "IFFCO Chowk, Sector 62 / Golf Course, Faridabad, Noida, Rohini - by name."),
        ("Pitched calls to customers OUTSIDE Delhi NCR", f'={k.count(PITCHED, ("Customer in Delhi NCR?", "N"))[0]}', "#,##0",
         "Lead city outside Delhi, Gurgaon, Noida, Greater Noida, Ghaziabad, Faridabad and nearby NCR towns - the hubs are out of reach."),
        ("Pitched calls where the customer raised a doubt", f'={k.count(PITCHED, ("Customer Doubt Category", "<>None"))[0]}',
         "#,##0", "At least one of the 6 doubt buckets, within the hospital part of the call."),
    ]
    for label, formula, fmt, how in headline:
        ws.cell(row=r, column=1, value=label)
        if label.startswith("Hospital pitch rate"):
            formula = f"=IFERROR({pitched_cell}/{analysed_cell},0)"
        ws.cell(row=r, column=2, value=formula)
        ws.cell(row=r, column=3, value=how)
        for c in range(1, 4):
            _body(ws.cell(row=r, column=c), fmt=fmt if c == 2 else None, bold=(c == 2))
        if label.startswith("Calls with a transcript"):
            analysed_cell = f"$B${r}"
        if label.startswith("Calls where the AGENT"):
            pitched_cell = f"$B${r}"
        r += 1

    r += 1
    _section(ws, r, "PITCH RATE BY SEARCH CLUSTER")
    r += 1
    _header(ws, r, ["Search cluster", "Calls analysed", "Agent pitched", "Pitch rate"])
    r += 1
    for cl, desc in [("brand", "Brand: pb health / pb hospital / policybazaar hospital / pb health hospital"),
                     ("hinglish", "Hinglish: khud ka / apna / humara hospital"),
                     ("hubs", "Hub names: iffco chowk, sector 62, golf course, faridabad, sector 2, sector 50, rohini")]:
        ws.cell(row=r, column=1, value=desc)
        ws.cell(row=r, column=2, value=f'={k.count(("Search Cluster", f"*{cl}*"), ("Transcript Analysed?", "Y"))[0]}')
        ws.cell(row=r, column=3, value=f'={k.count(("Search Cluster", f"*{cl}*"), PITCHED)[0]}')
        ws.cell(row=r, column=4, value=f"=IFERROR(C{r}/B{r},0)")
        for c in range(1, 5):
            _body(ws.cell(row=r, column=c), fmt="0.0%" if c == 4 else "#,##0" if c > 1 else None)
        r += 1

    r += 1
    _section(ws, r, "TOP 3 PITCH GAPS")
    r += 1
    rows = el["rows"]
    usps = ["Real Multispeciality Hospital", "Ethical Care / No Incentive to Overcharge", "Zero Waiting Time / Digital First",
            "PB Health App Download", "Exclusive Family Discount"]
    weakest = sorted(usps, key=lambda u: rows[u][0])[:2]
    hub_row, mg_row, none_row = rows["Any of the 5 hubs by name"][1], rows["MG Road, Gurgaon (address only)"][1], rows["No location named"][1]
    gap_lines = [
        f'="1. Hubs: only "&TEXT({E}!C{hub_row},"0.0%")&" of pitched calls name any of the 5 hubs ("&{E}!B{hub_row}&" calls). '
        f'Agents give the MG Road, Gurgaon address instead ("&TEXT({E}!C{mg_row},"0.0%")&") or no place at all ("&TEXT({E}!C{none_row},"0.0%")&")."',
    ]
    for i, u in enumerate(weakest, start=2):
        gap_lines.append(f'="{i}. {u}: said on "&TEXT({E}!C{rows[u][1]},"0.0%")&" of pitched calls ("&{E}!B{rows[u][1]}&" calls)."')
    for g in gap_lines:
        ws.cell(row=r, column=1, value=g)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        _body(ws.cell(row=r, column=1), bold=True)
        ws.row_dimensions[r].height = 30
        r += 1

    r += 1
    _section(ws, r, "HUB LOCATION MENTIONS (pitched calls)")
    r += 1
    _header(ws, r, ["Location named by the agent", "Calls", "% of pitched calls"])
    r += 1
    for label in ["IFFCO Chowk", "Sector 62 / Golf Course", "Faridabad Sec 2", "Noida Sec 50", "Rohini",
                  "Any of the 5 hubs by name", "MG Road, Gurgaon (address only)", "Gurgaon (generic)", "No location named"]:
        er = rows[label][1]
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=f"={E}!B{er}")
        ws.cell(row=r, column=3, value=f"={E}!C{er}")
        for c in range(1, 4):
            _body(ws.cell(row=r, column=c), fmt="0.0%" if c == 3 else "#,##0" if c == 2 else None,
                  bold=label.startswith("Any"))
        r += 1

    r += 1
    _section(ws, r, "CUSTOMER DOUBTS (pitched calls)")
    r += 1
    _header(ws, r, ["Doubt", "Calls", "% of pitched calls"])
    r += 1
    D = f"'{DOUBTS}'"
    for d in doubt_rows:
        ws.cell(row=r, column=1, value=f"{d['label']} - {cf.DOUBT_DESC[d['key']]}")
        ws.cell(row=r, column=2, value=f"={D}!C{d['row']}")
        ws.cell(row=r, column=3, value=f"={D}!D{d['row']}")
        for c in range(1, 4):
            _body(ws.cell(row=r, column=c), fmt="0.0%" if c == 3 else "#,##0" if c == 2 else None)
        r += 1

    r += 1
    _section(ws, r, "HOW TO READ THIS")
    r += 1
    notes = [
        "Every number is a count of real calls. Filter the 'Call-by-Call Master Data' tab to see the calls behind any figure.",
        "'Pitched' = the agent named Policybazaar's OWN hospital. 'PB Health' alone is not enough: on real calls it is mostly "
        "the PB Health app, 'PB Health wellness' or platinum membership.",
        "Story elements, hubs and doubts are only read in the hospital part of the call (the mention and the few turns after "
        "it), so policy talk about cashless or discounts elsewhere in the call is not counted.",
        "Transcripts are the portal's romanised Hinglish speech-to-text. Spelling variants are normalised (khud ke/ki -> khud "
        "ka, policy bazar -> policybazaar, sec 62 -> sector 62, 'baby health hospital' -> pb health hospital).",
        "Detection is keyword/phrase based and was tuned by reading matched lines on real calls: about 8 in 10 detected doubts "
        "are genuine hospital doubts. In a read of 45 pitched calls with no doubt detected, customers mostly just acknowledged "
        "('theek hai', 'ji ji') or went back to the premium, so a low doubt rate is real - customers rarely engage with the pitch. "
        "Treat doubt counts as a floor.",
        *(["The last day in the period is today, which the portal is still indexing, so its numbers are incomplete."]
          if df["Call Date"].max() >= date.today() else []),
        ("Search coverage: every search window was read in full." if not gaps else
         f"Search coverage: {len(gaps)} search window(s) never answered (portal timeouts), so calls in them are missing: "
         + "; ".join(f"{c} {a[:16]} to {b[11:16]}" for c, a, b in gaps) + "."),
    ]
    for n in notes:
        ws.cell(row=r, column=1, value=f"• {n}")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        _body(ws.cell(row=r, column=1))
        ws.row_dimensions[r].height = 30
        r += 1
    for col, w in {1: 62, 2: 16, 3: 70, 4: 14}.items():
        ws.column_dimensions[get_column_letter(col)].width = w


# ---- Tab 4 -----------------------------------------------------------------
def _style_master(ws, df: pd.DataFrame) -> None:
    widths = {"Call ID": 13, "Call Date": 11, "Talktime (mins)": 9, "Agent ID": 11, "Customer City": 14,
              "Pitched? (Y/N)": 9, "Hubs Mentioned": 28, "Multispeciality?": 22, "No Overcharge Pitch?": 10,
              "App Pushed?": 9, "Customer Doubt Category": 26, "Customer Quote (Verbatim)": 60,
              "Agent Quote (Verbatim)": 70, "Other Agent Claims": 30, "Pitch Phrase Used": 34, "Search Cluster": 16}
    for i, col in enumerate(MASTER_COLS, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = widths.get(col, 12)
        h = ws.cell(row=1, column=i)
        h.font = _font(bold=True, color="FFFFFF")
        h.fill = PatternFill("solid", fgColor=NAVY if i <= 13 else GREY)
        h.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 42
    date_col = MASTER_COLS.index("Call Date") + 1
    id_col = MASTER_COLS.index("Call ID") + 1
    for row in ws.iter_rows(min_row=2, max_row=len(df) + 1):
        for c in row:
            c.font = _font(size=9)
        row[date_col - 1].number_format = "yyyy-mm-dd"
        row[id_col - 1].number_format = "0"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(MASTER_COLS))}{len(df) + 1}"


# ---- console ---------------------------------------------------------------
def _print_console(k: Counter, df: pd.DataFrame, period: str, doubt_rows: list[dict], out: Path,
                   gaps: list[tuple[str, str, str]]) -> None:
    scanned = len(df)
    analysed = k.count(("Transcript Analysed?", "Y"))[1]
    came_up = sum(k.count(("Hospital Raised By", v))[1] for v in ("Agent", "Customer", "Both"))
    pitched = k.count(PITCHED)[1]
    pbh_only = k.count((PBH_COL, "Y"), ("Pitched? (Y/N)", "N"))[1]
    hubs5 = k.count(PITCHED, ("Named One of 5 Hubs?", "Y"))[1]
    outside = k.count(PITCHED, ("Customer in Delhi NCR?", "N"))[1]
    any_doubt = k.count(PITCHED, ("Customer Doubt Category", "<>None"))[1]
    pct = lambda a, b: f"{(100 * a / b if b else 0):5.1f}%"
    line = "=" * 78
    print(f"\n{line}\n  PB HEALTH HOSPITAL - REAL CALL ANALYSIS   ({period})\n{line}")
    print(f"  Calls scanned (search hits) ..................... {scanned:>7,}")
    print(f"  Calls with transcript analysed .................. {analysed:>7,}")
    print(f"  Calls where PB Hospital came up (either side) ... {came_up:>7,}")
    print(f"  Calls where the AGENT pitched PB Hospital ....... {pitched:>7,}   {pct(pitched, analysed)} of analysed")
    print(f"  Agent said 'PB Health' but not the hospital ..... {pbh_only:>7,}")
    print(f"  Pitched calls naming one of the 5 hubs .......... {hubs5:>7,}   {pct(hubs5, pitched)} of pitched")
    print(f"  Pitched calls to customers outside Delhi NCR .... {outside:>7,}   {pct(outside, pitched)} of pitched")
    print(f"  Pitched calls with a customer doubt ............. {any_doubt:>7,}   {pct(any_doubt, pitched)} of pitched")
    print(f"\n  TOP CUSTOMER DOUBTS (base: {pitched:,} pitched calls)")
    print(f"  {'Doubt':<30}{'Calls':>8}{'% pitched':>12}")
    print(f"  {'-' * 50}")
    for d in doubt_rows:
        print(f"  {d['label']:<30}{d['n']:>8,}{pct(d['n'], pitched):>12}")
    if gaps:
        print(f"\n  WARNING: {len(gaps)} search window(s) timed out and are missing:")
        for c, a, b in gaps:
            print(f"    {c:9} {a} -> {b}")
    print(f"\n  Excel saved to: {out.resolve()}\n{line}\n")
