"""End-to-end test of export_call_analysis with a fake portal client (no network)."""
from datetime import date

import openpyxl
import pytest

import api_client
import export_call_analysis as exp

PITCH_CALL = [
    ("agent", "Sir Policybazaar ka khud ka multispeciality hospital hai, Noida Sector 50 mein."),
    ("customer", "Wahan surgery bhi hoti hai kya?"),
    ("agent", "Haan sir, OT ICU sab hai, doctors salary pe hain."),
    ("customer", "Meri policy cashless chalegi?"),
    ("agent", "Main check karke batata hoon."),
]
CLINIC_CALL = [
    ("agent", "PB Health ka OPD pass milega aapko, Gurgaon mein clinic hai."),
    ("customer", "Not interested, aage chalo."),
    ("agent", "Theek hai sir."),
]
NO_PITCH_CALL = [("agent", "Aap Faridabad se ho?"), ("customer", "Haan ji")]
TRANSCRIPTS = {"101": PITCH_CALL, "102": CLINIC_CALL, "103": NO_PITCH_CALL, "104": []}


class FakeClient:
    searches: list = []

    def __init__(self, *a, **k):
        pass

    def login(self):
        pass

    def search_window(self, terms, w_from, w_to, categories=None):
        FakeClient.searches.append((tuple(terms), categories))
        if not w_from.date() <= date(2026, 10, 14) <= w_to.date():
            return [], 0
        ids = {"pb health": ["101", "102"], "faridabad": ["103"], "": ["101", "104"]}.get(terms[0] if terms else "", [])
        rows = [{"_source": {"callDataId": i, "recordingCallDate": "2026-10-14 11:00:00", "agentId": f"A{i}",
                             "lead_city_name": "Noida", "talktime": 300}} for i in ids]
        return rows, len(rows)

    def get_transcript(self, call_id):
        if call_id == "999":
            raise RuntimeError("boom")
        return {"english": [{"speakerId": s, "text": t} for s, t in TRANSCRIPTS[call_id]], "hindi": []}

    def count_qualified_ncr_calls(self, day):
        return (100, "api_total") if day == date(2026, 10, 14) else (None, "unavailable")


@pytest.fixture()
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(api_client, "PortalClient", FakeClient)
    monkeypatch.setattr(exp, "CACHE_DIR", tmp_path / "cache")
    FakeClient.searches = []
    return tmp_path


def _sheet(path, name):
    ws = openpyxl.load_workbook(path)[name]
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_live_run_builds_four_tab_workbook(env, capsys):
    out = env / "PB_Health_Call_Analysis.xlsx"
    assert exp.main(["--from", "2026-10-13", "--to", "2026-10-14", "--out", str(out)]) == 0

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Executive Summary", "Customer Doubts & Objections",
                             "Pitch Elements Tracker", "Call-by-Call Master Data"]
    # every search phrase + the category-only query, per 2-day window
    assert {t[0] for t, _ in FakeClient.searches if t} == {t for v in exp.SEARCH_TERMS.values() for t in v}
    assert ((), [10693]) in FakeClient.searches

    master = _sheet(out, "Call-by-Call Master Data")
    assert master[0][:13] == ["Call ID", "Call Date", "Talktime (mins)", "Agent ID", "Customer City", "Pitched? (Y/N)",
                              "Hubs Mentioned", "Multispeciality?", "No Overcharge Pitch?", "App Pushed?",
                              "Customer Doubt Category", "Customer Quote (Verbatim)", "Agent Quote (Verbatim)"]
    rows = {r[0]: dict(zip(master[0], r)) for r in master[1:]}
    assert set(rows) == {"101", "102", "103"}                       # 104 had an empty transcript
    assert rows["101"]["Pitched? (Y/N)"] == "Y" and rows["101"]["Hubs Mentioned"] == "Noida Sec 50"
    assert rows["101"]["Multispeciality?"] == "Yes" and rows["101"]["No Overcharge Pitch?"] == "Yes"
    assert rows["101"]["Customer Doubt Category"] == "Scope doubt; Cashless/Insurance doubt"
    assert "[Scope doubt] Wahan surgery bhi hoti hai kya?" in rows["101"]["Customer Quote (Verbatim)"]
    assert rows["101"]["Talktime (mins)"] == 5 and rows["101"]["Search Terms Matched"] == "category 10693, pb health"
    assert rows["102"]["Multispeciality?"] == "No - called it clinic/OPD"
    assert rows["102"]["Hubs Mentioned"] == "Gurgaon (generic)"
    assert rows["103"]["Pitched? (Y/N)"] == "N" and rows["103"]["Multispeciality?"] == "N/A"

    doubts = _sheet(out, "Customer Doubts & Objections")
    assert doubts[0] == ["Doubt Category", "What Customer Actually Asks (Verbatim Quotes)", "Frequency (Call Count)",
                         "% of Pitched Calls", "How Agents Currently Answer (Real Example)",
                         "Recommended Counter-Script for Sales Deck"]
    freq = {r[0]: (r[2], r[3]) for r in doubts[1:]}
    assert freq["Scope doubt"] == (1, 0.5) and freq["Skepticism / Rushing past"] == (1, 0.5)
    assert freq["Discount confusion"] == (0, 0.0)
    assert [r[2] for r in doubts[1:]] == sorted([r[2] for r in doubts[1:]], reverse=True)

    summary = {r[0]: r[1] for r in _sheet(out, "Executive Summary") if r[0]}
    assert summary["Calls scanned (unique calls returned by search)"] == 4
    assert summary["Calls analyzed (full transcript fetched)"] == 3
    assert summary["Calls where the agent pitched PB Hospital"] == 2
    assert summary["Hospital pitch rate (pitched / analyzed)"] == pytest.approx(2 / 3)
    assert summary["All qualified NCR health sales calls in window (portal count)"] == 100

    elements = {r[0]: tuple(r[1:4]) for r in _sheet(out, "Pitch Elements Tracker")[3:]}
    assert elements["Hub named: Noida Sec 50"] == (1, 0.5, 1)
    assert elements['Called it a "Clinic" / "OPD pass" (wrong framing)'][0] == 1

    printed = capsys.readouterr().out
    assert "TOP CUSTOMER DOUBTS" in printed and str(out.resolve()) in printed
    assert (env / "PB_Health_Harmonized_Dialogue.csv").exists()


def test_cached_rerun_needs_no_api(env, monkeypatch):
    out = env / "a.xlsx"
    exp.main(["--from", "2026-10-13", "--to", "2026-10-14", "--out", str(out), "--no-denominator"])

    class Boom(FakeClient):
        def __init__(self, *a, **k):
            raise AssertionError("API must not be used")

    monkeypatch.setattr(api_client, "PortalClient", Boom)
    out2 = env / "b.xlsx"
    assert exp.main(["--from", "2026-10-13", "--to", "2026-10-14", "--cached", "--out", str(out2)]) == 0
    assert _sheet(out, "Call-by-Call Master Data") == _sheet(out2, "Call-by-Call Master Data")


def test_failed_transcripts_are_reported(env):
    def search(self, terms, w_from, w_to, categories=None):
        return [{"callDataId": "999", "recordingCallDate": "2026-10-14"}, {"callDataId": "101"}], 2
    FakeClient.search_window = search
    try:
        out = env / "c.xlsx"
        assert exp.main(["--days", "1", "--out", str(out), "--no-denominator", "--no-category"]) == 0
        summary = {r[0]: r[1] for r in _sheet(out, "Executive Summary") if r[0]}
        assert summary["Calls skipped (transcript failed or empty)"] == 1
    finally:
        del FakeClient.search_window
